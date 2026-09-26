#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""Create a teacher-facing DOCX report for GINSFGO_LC vs GREAT-MSF LC."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULT_DIR = ROOT / "plot" / "results" / "GINSFGO_LC_vs_MSF（2）"
FIGURE_DIR = RESULT_DIR / "ENU_error_series"
SUMMARY_CSV = RESULT_DIR / "comparison_summary.csv"
PAIRWISE_CSV = RESULT_DIR / "comparison_pairwise.csv"
OUTPUT_DOCX = RESULT_DIR / "GREAT_FGO与GREAT_MSF松耦合定位精度对比报告.docx"

NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = RGBColor(0, 0, 0)


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = TEXT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_repeat_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        set_run_font(run, east_asia="微软雅黑", latin="Arial", size=16 if level == 1 else 13, bold=True)
    return paragraph


def add_body(doc, text, bold_lead=None, indent=True):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(6)
    if indent:
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.3
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5)
    return paragraph


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    set_table_borders(table)

    header = table.rows[0]
    repeat_table_header(header)
    set_repeat_row_no_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=110, bottom=110, start=90, end=90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(value))
        set_run_font(run, east_asia="微软雅黑", latin="Arial", size=font_size, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if widths:
            cell.width = Cm(widths[index])

    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_repeat_row_no_split(row)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell, top=95, bottom=95, start=85, end=85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size)
            if widths:
                cell.width = Cm(widths[col_index])

    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure_page(doc, figure_number, case, image_path, interpretation):
    doc.add_page_break()
    add_heading(doc, f"{figure_number + 4} {case} ENU位置误差序列", level=1)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.3))
    add_caption(doc, f"图{figure_number} {case}中GINSFGO松耦合与MSF松耦合的ENU位置误差序列")
    add_body(doc, interpretation)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = TEXT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("GREAT FGO与GREAT MSF松耦合定位精度对比报告")
    set_run_font(run, east_asia="微软雅黑", latin="Arial", size=8.5)
    run.font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])


def main():
    for required in (SUMMARY_CSV, PAIRWISE_CSV):
        if not required.is_file():
            raise FileNotFoundError(required)

    summary = pd.read_csv(SUMMARY_CSV)
    pairwise = pd.read_csv(PAIRWISE_CSV)
    methods = summary.set_index(["case", "method"])

    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    title.paragraph_format.space_after = Pt(16)
    run = title.add_run("GREAT FGO与GREAT MSF松耦合定位精度对比报告")
    set_run_font(run, east_asia="微软雅黑", latin="Arial", size=24, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    run = subtitle.add_run("基于四组实测GNSS INS数据的共同历元评定")
    set_run_font(run, east_asia="微软雅黑", latin="Arial", size=14)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.paragraph_format.space_before = Pt(190)
    run = date_p.add_run("2026年9月")
    set_run_font(run, size=12)
    doc.add_page_break()

    add_heading(doc, "摘要", level=1)
    add_body(
        doc,
        "本报告比较GINSFGO松耦合因子图优化与GREAT MSF松耦合滤波在四组实测数据上的定位性能。为避免输出起止时间和固定状态差异造成偏差，所有指标均在两种方法共同存在的历元上计算，并使用同一IMU中心真值。实测结果表明，GINSFGO在Data06、20211012和20250928三个算例上获得更低的三维位置RMSE，MSF在Data05上更好。因此，当前结果支持因子图优化在多数场景中提高位置精度和异常抑制能力的判断，但不支持其在所有数据和全部导航指标上必然占优。",
    )
    add_body(
        doc,
        "核心结论是：因子图优化具有重复线性化和多历元联合约束的理论优势，当前实现的定位结果也在四个算例中的三个表现更好；其优势仍受滑动窗口长度、因子权重、初始状态和异常值处理影响。",
        bold_lead="核心结论是：",
    )

    summary_rows = []
    for row in pairwise.itertuples(index=False):
        improvement = (row.msf_pos_3d_rmse_m - row.ginsfgo_pos_3d_rmse_m) / row.msf_pos_3d_rmse_m * 100.0
        verdict = "FGO更好" if improvement > 0 else "MSF更好"
        summary_rows.append(
            [
                row.case,
                f"{int(row.common_epochs)}",
                f"{row.ginsfgo_pos_3d_rmse_m:.3f}",
                f"{row.msf_pos_3d_rmse_m:.3f}",
                f"{improvement:+.1f}%",
                verdict,
            ]
        )
    add_caption(doc, "表1 四个算例的三维位置RMSE总体比较")
    add_table(
        doc,
        ["算例", "共同历元", "FGO RMSE m", "MSF RMSE m", "FGO相对变化", "结果"],
        summary_rows,
        widths=[2.2, 2.0, 2.4, 2.4, 2.4, 2.1],
        font_size=9,
    )

    add_heading(doc, "1 研究目的与比较对象", level=1)
    add_body(
        doc,
        "本次比较关注两种GNSS INS松耦合实现。GREAT MSF采用递推滤波框架，在每个历元利用GNSS位置解更新惯性导航状态；GINSFGO采用滑动窗口因子图，将GNSS位置因子与IMU预积分因子放入同一优化问题。两种方法输出的INS位置均对应IMU中心，因此可使用相同真值进行比较。",
    )
    add_body(
        doc,
        "比较包含Data05车辆开阔场景、Data06无人机开阔场景、20211012城市车辆场景和20250928校园场景。四组结果均使用关闭自主对准并给定有效位置、速度和姿态初值的配置，以减少对准时段不同造成的样本选择偏差。",
    )

    add_heading(doc, "2 理论分析", level=1)
    add_heading(doc, "2 1 MSF松耦合滤波", level=2)
    add_body(
        doc,
        "MSF松耦合在每个历元将GNSS位置作为量测输入卡尔曼滤波器。该方法计算量小、实时性强，并且在模型和噪声参数合适时可以获得稳定结果。其主要限制是状态在当前线性化点附近递推更新，历史量测通常被压缩进当前状态及协方差，难以重新调整早期线性化误差。",
    )
    add_heading(doc, "2 2 GINSFGO松耦合因子图", level=2)
    add_body(
        doc,
        "GINSFGO使用GNSS位置因子、IMU预积分因子和边缘化先验共同约束滑动窗口内的状态。优化器可以在窗口内反复线性化并联合修正多个历元，因此在运动非线性较强、GNSS解波动或单历元异常较明显时，理论上具有更好的位置估计潜力。",
    )
    add_body(
        doc,
        "这种优势不是无条件的。当前配置中的滑动窗口较短，求解时间、最大迭代次数、因子权重和鲁棒核都会限制优化效果。因子图还依赖GNSS位置因子的质量；如果输入位置已经含有系统误差，优化器不能仅凭计算形式消除该误差。",
    )

    add_heading(doc, "3 数据与评定方法", level=1)
    add_body(
        doc,
        "评定首先将GINSFGO与MSF结果分别按GPS周内秒匹配到对应IMU中心真值，匹配容差为0.05 s；随后仅保留两种方法共同存在的真值历元。四组数据的最大实际时间差不超过0.009998 s。位置误差由ECEF坐标差转换到固定参考点下的东、北、天方向。三维位置RMSE按各历元东、北、天误差平方和的均值开方计算。",
    )
    add_body(
        doc,
        "除全体共同历元外，报告还检查了双方均为Fixed状态的共同子集、位置误差P95和最大误差。Fixed率只在共同输出历元中计算，用于描述模糊度状态，不作为因子图位置优化效果的唯一指标。",
    )
    method_rows = []
    truth_map = {
        "Data05": "groundtruth_HG4930",
        "Data06": "groundtruth_HG4930",
        "20211012": "groundtruth_1012_ADIS",
        "20250928": "groundtruth_0928_MTI",
    }
    for row in pairwise.itertuples(index=False):
        method_rows.append([row.case, int(row.common_epochs), int(row.both_fixed_epochs), truth_map[row.case]])
    add_caption(doc, "表2 共同历元与参考真值")
    add_table(doc, ["算例", "共同历元", "共同Fixed历元", "IMU中心真值"], method_rows, widths=[2.5, 2.6, 3.0, 5.4], font_size=9)

    doc.add_page_break()
    add_heading(doc, "4 定位精度结果", level=1)
    position_rows = []
    for row in pairwise.itertuples(index=False):
        g = methods.loc[(row.case, "GINSFGO_LC")]
        m = methods.loc[(row.case, "MSF_LC")]
        position_rows.append(
            [
                row.case,
                f"{g.pos_3d_rmse_m:.3f}",
                f"{m.pos_3d_rmse_m:.3f}",
                f"{g.pos_3d_p95_m:.3f}/{m.pos_3d_p95_m:.3f}",
                f"{g.pos_3d_max_m:.3f}/{m.pos_3d_max_m:.3f}",
                f"{g.fixed_rate_pct:.1f}%/{m.fixed_rate_pct:.1f}%",
            ]
        )
    add_caption(doc, "表3 位置精度与Fixed率比较")
    add_table(
        doc,
        ["算例", "FGO 3D RMSE m", "MSF 3D RMSE m", "P95 FGO/MSF m", "最大值 FGO/MSF m", "Fixed率 FGO/MSF"],
        position_rows,
        widths=[2.0, 2.5, 2.5, 2.8, 3.1, 3.1],
        font_size=8.3,
    )
    add_body(
        doc,
        "Data05中MSF的三维位置RMSE为0.063 m，优于GINSFGO的0.077 m，说明在开阔、观测质量较好且动态较平稳的车辆场景中，递推滤波已经可以获得很高的精度，因子图的短窗口优化没有形成额外收益。",
    )
    add_body(
        doc,
        "Data06、20211012和20250928中，GINSFGO的三维位置RMSE分别比MSF降低53.2%、14.5%和57.4%。其中20250928的MSF在末段出现86.533 m的三维位置异常，主要来自天向误差；GINSFGO也出现异常，但最大值为19.260 m，表明因子图在该算例中减弱了异常量测的影响。",
    )
    add_body(
        doc,
        "Data06需要结合P95理解：MSF的P95为0.469 m，略优于GINSFGO的0.505 m，但MSF存在5.653 m的早期位置尖峰，因此全局RMSE更差。这说明GINSFGO在该算例中的主要优势来自对少数大误差的抑制，而不是所有时段均保持更小误差。",
    )

    figure_interpretations = {
        "Data05": "Data05中两种方法长期保持厘米级误差。MSF的北向和天向RMSE略低，GINSFGO在少数历元出现更明显的负向尖峰，因此MSF取得更低的三维位置RMSE。",
        "Data06": "Data06的主要差异集中在开始约300 s。MSF在北向出现约5 m的早期尖峰，GINSFGO也有初始化波动但幅度较小；稳定后两条曲线均接近零。",
        "20211012": "20211012中两种方法在多个遮挡或观测质量下降时段出现同步波动。GINSFGO的天向RMSE低于MSF，而MSF的北向RMSE较低，最终GINSFGO取得更低的三维位置RMSE。",
        "20250928": "20250928的末段出现显著异常。MSF天向误差峰值超过80 m，GINSFGO峰值约19 m。因子图没有完全消除异常，但显著降低了最大误差及总体三维RMSE。",
    }
    for number, case in enumerate(["Data05", "Data06", "20211012", "20250928"], start=1):
        add_figure_page(
            doc,
            number,
            case,
            FIGURE_DIR / f"{case}_ENU_error_series.png",
            figure_interpretations[case],
        )

    doc.add_page_break()
    add_heading(doc, "9 速度姿态与Fixed率结果", level=1)
    support_rows = []
    for row in pairwise.itertuples(index=False):
        g = methods.loc[(row.case, "GINSFGO_LC")]
        m = methods.loc[(row.case, "MSF_LC")]
        support_rows.append(
            [
                row.case,
                f"{g.vel_3d_rmse_mps:.3f}/{m.vel_3d_rmse_mps:.3f}",
                f"{g.heading_rmse_deg:.3f}/{m.heading_rmse_deg:.3f}",
                f"{g.pitch_rmse_deg:.3f}/{m.pitch_rmse_deg:.3f}",
                f"{g.roll_rmse_deg:.3f}/{m.roll_rmse_deg:.3f}",
            ]
        )
    add_caption(doc, "表4 速度与姿态RMSE比较 数值顺序为FGO MSF")
    add_table(
        doc,
        ["算例", "三维速度 m/s", "航向 °", "俯仰 °", "横滚 °"],
        support_rows,
        widths=[2.2, 3.4, 3.0, 3.0, 3.0],
        font_size=8.8,
    )
    add_body(
        doc,
        "MSF在四个算例中的航向RMSE均低于GINSFGO，并在前三个算例中取得更高Fixed率。速度和俯仰横滚结果呈混合状态，没有一种方法稳定占优。因此，当前因子图实现的优势主要体现在位置估计，尚不能概括为完整导航状态精度全面提高。",
    )

    add_heading(doc, "10 结果讨论", level=1)
    add_body(
        doc,
        "实测结果总体符合因子图在位置优化上的理论预期。GINSFGO在三个算例中获得更低位置RMSE，在异常最明显的20250928中也显著压低了最大误差。这与因子图能够联合利用多个历元并重新线性化历史状态的特点一致。",
    )
    add_body(
        doc,
        "Data05的反例说明该优势并非由求解框架自动保证。当前滑动窗口长度较短，且GNSS位置因子、IMU噪声、鲁棒核和边缘化先验尚未针对不同平台分别调优。在观测条件好、滤波模型匹配的场景中，MSF可能达到相同或更好的结果。",
    )
    add_body(
        doc,
        "本报告不对四个算例的RMSE直接求平均，因为各算例持续时间、平台动态和误差量级不同。采用胜出算例数量、逐算例RMSE、P95和最大误差共同判断，可以避免大误差场景支配总体结论。",
    )

    add_heading(doc, "11 结论", level=1)
    add_body(
        doc,
        "在统一初始状态、同一IMU中心真值和共同历元条件下，GINSFGO松耦合在四个算例中的三个取得更低三维位置RMSE。其在Data06、20211012和20250928上的RMSE相对MSF分别降低53.2%、14.5%和57.4%，支持因子图优化具有更好总体定位精度与异常抑制潜力的判断。",
    )
    add_body(
        doc,
        "同时，MSF在Data05中表现更好，并在航向和Fixed率方面具有明显优势。适合向老师汇报的结论应表述为：当前GINSFGO实现的位置精度总体优于MSF，但尚未在速度、姿态和所有数据场景中形成全面优势。后续应重点增加滑动窗口长度敏感性试验，并针对Data06早期尖峰和20250928末段异常检查因子权重及鲁棒核。",
    )

    add_heading(doc, "附录 复现说明", level=1)
    add_bullet(doc, "精度比较脚本：plot/compare_ginsfgo_lc_msf.py")
    add_bullet(doc, "ENU误差绘图脚本：plot/plot_lc_enu_error_series.py")
    add_bullet(doc, "汇总指标：plot/results/GINSFGO_LC_vs_MSF（2）/comparison_summary.csv")
    add_bullet(doc, "配对比较：plot/results/GINSFGO_LC_vs_MSF（2）/comparison_pairwise.csv")
    add_bullet(doc, "所有图表均使用GINSFGO与MSF共同历元及同一IMU中心真值。")

    doc.core_properties.title = "GREAT FGO与GREAT MSF松耦合定位精度对比报告"
    doc.core_properties.subject = "四组GNSS INS实测数据的共同历元精度比较"
    doc.core_properties.keywords = "GREAT FGO, GREAT MSF, GNSS INS, 松耦合, 因子图, ENU"
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
